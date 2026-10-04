from pathlib import Path
import json
import re
import logging

from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
from docx.table import Table

import PyPDF2
import openpyxl

logger = logging.getLogger(__name__)

class FileProcessor:
    SUPPORTED_EXTENSIONS = {
        '.txt': 'text',
        '.md': 'text',
        '.docx': 'docx',
        '.pdf': 'pdf',
        '.xlsx': 'excel',
        '.xls': 'excel',
        '.json': 'json',
        '.html': 'html'
    }

    @staticmethod
    def extract_text(file_path: Path) -> str:
        ext = file_path.suffix.lower()

        if ext not in FileProcessor.SUPPORTED_EXTENSIONS:
            logger.warning(f"Неподдерживаемый формат файла: {ext}")
            return f"[Файл {file_path.name} - неподдерживаемый формат]"

        try:
            if ext in ['.txt', '.md']:
                return FileProcessor._read_text_file(file_path)
            elif ext == '.docx':
                return FileProcessor._read_docx_file(file_path)
            elif ext == '.pdf':
                return FileProcessor._read_pdf_file(file_path)
            elif ext in ['.xlsx', '.xls']:
                return FileProcessor._read_excel_file(file_path)
            elif ext == '.json':
                return FileProcessor._read_json_file(file_path)
            elif ext == '.html':
                return FileProcessor._read_html_file(file_path)
            else:
                return f"[Файл {file_path.name} - содержимое не может быть извлечено]"
        except Exception as e:
            logger.error(f"Ошибка при чтении файла {file_path}: {str(e)}")
            return f"[Ошибка чтения файла {file_path.name}: {str(e)}]"

    @staticmethod
    def _read_text_file(file_path: Path) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()

    @staticmethod
    def _iter_block_items(parent):
        parent_elm = parent.element.body
        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    @staticmethod
    def _table_to_text(table):
        lines = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
        return "\n".join(lines)

    @staticmethod
    def _read_docx_file(file_path: Path) -> str:
        doc = DocxDocument(file_path)
        full_text = []

        for block in FileProcessor._iter_block_items(doc):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if text:
                    full_text.append(text)
            elif isinstance(block, Table):
                full_text.append("[TABLE]")
                full_text.append(FileProcessor._table_to_text(block))
                full_text.append("")

        return "\n".join(full_text)

    @staticmethod
    def _read_pdf_file(file_path: Path) -> str:
        text = []
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text.append(page.extract_text())
        return '\n'.join(text)

    @staticmethod
    def _read_excel_file(file_path: Path) -> str:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        text_parts = []

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text_parts.append(f"--- Лист: {sheet_name} ---")
            for row in sheet.iter_rows(values_only=True):
                row_text = ' | '.join(str(cell) for cell in row if cell is not None)
                if row_text:
                    text_parts.append(row_text)

        return '\n'.join(text_parts)

    @staticmethod
    def _read_json_file(file_path: Path) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return json.dumps(data, ensure_ascii=False, indent=2)

    @staticmethod
    def _read_html_file(file_path: Path) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        text = re.sub(r'<[^>]+>', ' ', content)
        text = re.sub(r'\s+', ' ', text)
        return text.strip()
